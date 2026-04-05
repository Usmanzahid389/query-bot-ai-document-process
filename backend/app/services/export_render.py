"""Build PDF/DOCX bytes for export (no external binaries required; uses reportlab + python-docx)."""

from __future__ import annotations

import io
from typing import Literal

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

Format = Literal["pdf", "docx"]


def render_pdf(title: str, lines: list[str]) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    width, height = letter
    y = height - 50
    c.setFont("Helvetica-Bold", 14)
    for part in title.split("\n"):
        c.drawString(50, y, part[:120])
        y -= 18
    y -= 10
    c.setFont("Helvetica", 10)
    for para in lines:
        for line in _wrap(para, 90):
            if y < 50:
                c.showPage()
                y = height - 50
                c.setFont("Helvetica", 10)
            c.drawString(50, y, line)
            y -= 12
        y -= 6
    c.save()
    return buf.getvalue()


def _wrap(text: str, width: int) -> list[str]:
    words = text.replace("\r", "").split()
    if not words:
        return [""]
    lines: list[str] = []
    cur: list[str] = []
    n = 0
    for w in words:
        add = len(w) + (1 if cur else 0)
        if n + add > width and cur:
            lines.append(" ".join(cur))
            cur = [w]
            n = len(w)
        else:
            cur.append(w)
            n += add
    if cur:
        lines.append(" ".join(cur))
    return lines


def render_docx(title: str, lines: list[str]) -> bytes:
    doc = DocxDocument()
    doc.add_heading(title.split("\n")[0][:200], level=1)
    if "\n" in title:
        doc.add_paragraph(title.split("\n", 1)[1])
    for para in lines:
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
